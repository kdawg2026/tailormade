"""DOCX and PDF rendering for the tailored resume and cover letter."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


# ---------------------------------------------------------------- DOCX ----

def _docx_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # Bottom border under the heading, matching the original resume style.
    from docx.oxml.ns import qn

    p_pr = p._p.get_or_add_pPr()
    borders = p_pr.makeelement(qn("w:pBdr"), {})
    bottom = p_pr.makeelement(
        qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "999999"},
    )
    borders.append(bottom)
    p_pr.append(borders)


def _docx_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)


def build_docx(resume: dict, path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(54)
        section.left_margin = section.right_margin = Pt(54)

    style = doc.styles["Normal"]
    style.font.name = "Cambria"
    style.font.size = Pt(10.5)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(resume["name"])
    name_run.bold = True
    name_run.font.size = Pt(20)

    for line in (resume["tagline"], resume["contact"]):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)

    _docx_heading(doc, "PROFESSIONAL SUMMARY")
    doc.add_paragraph(resume["summary"])

    _docx_heading(doc, "CORE COMPETENCIES")
    for item in resume["competencies"]:
        _docx_bullet(doc, item)

    _docx_heading(doc, "PROFESSIONAL EXPERIENCE")
    for job in resume["experience"]:
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(6)
        title_p.paragraph_format.space_after = Pt(0)
        title_p.add_run(job["title"]).bold = True
        org_p = doc.add_paragraph()
        org_p.paragraph_format.space_after = Pt(2)
        org_p.add_run(job["org_line"]).bold = True
        for bullet in job["bullets"]:
            _docx_bullet(doc, bullet)

    _docx_heading(doc, "CERTIFICATIONS")
    for cert in resume["certifications"]:
        _docx_bullet(doc, cert)

    _docx_heading(doc, "EDUCATION")
    for entry in resume["education"]:
        p = doc.add_paragraph()
        p.add_run(entry).bold = True

    _docx_heading(doc, "PROJECTS")
    for project in resume["projects"]:
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(4)
        title_p.paragraph_format.space_after = Pt(0)
        title_p.add_run(project["title"]).bold = True
        doc.add_paragraph(project["description"])

    doc.save(path)


# ----------------------------------------------------------------- PDF ----

def build_pdf(resume: dict, path: Path) -> None:
    body = ParagraphStyle("body", fontName="Times-Roman", fontSize=10.5, leading=13.5)
    name_style = ParagraphStyle(
        "name", parent=body, fontName="Times-Bold", fontSize=20, leading=24, alignment=1
    )
    center = ParagraphStyle("center", parent=body, alignment=1)
    heading = ParagraphStyle(
        "heading", parent=body, fontName="Times-Bold", fontSize=12,
        leading=14, spaceBefore=10, spaceAfter=1,
    )
    job_title = ParagraphStyle("jobTitle", parent=body, fontName="Times-Bold", spaceBefore=6)
    job_org = ParagraphStyle("jobOrg", parent=body, fontName="Times-Bold")

    def esc(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def bullets(items: list[str]) -> ListFlowable:
        return ListFlowable(
            [ListItem(Paragraph(esc(i), body), leftIndent=18) for i in items],
            bulletType="bullet", bulletFontSize=8, start="•", spaceBefore=2,
        )

    def section(title: str) -> list:
        return [
            Paragraph(esc(title), heading),
            HRFlowable(width="100%", thickness=0.7, color="#999999", spaceAfter=4),
        ]

    story = [
        Paragraph(esc(resume["name"]), name_style),
        Spacer(1, 4),
        Paragraph(esc(resume["tagline"]), center),
        Paragraph(esc(resume["contact"]), center),
        *section("PROFESSIONAL SUMMARY"),
        Paragraph(esc(resume["summary"]), body),
        *section("CORE COMPETENCIES"),
        bullets(resume["competencies"]),
        *section("PROFESSIONAL EXPERIENCE"),
    ]
    for job in resume["experience"]:
        story.append(Paragraph(esc(job["title"]), job_title))
        story.append(Paragraph(esc(job["org_line"]), job_org))
        story.append(bullets(job["bullets"]))

    story += section("CERTIFICATIONS") + [bullets(resume["certifications"])]
    story += section("EDUCATION")
    story += [Paragraph(esc(e), job_org) for e in resume["education"]]
    story += section("PROJECTS")
    for project in resume["projects"]:
        story.append(Paragraph(esc(project["title"]), job_title))
        story.append(Paragraph(esc(project["description"]), body))

    SimpleDocTemplate(
        str(path), pagesize=letter,
        leftMargin=0.75 * inch, rightMargin=0.75 * inch,
        topMargin=0.75 * inch, bottomMargin=0.75 * inch,
    ).build(story)


# --------------------------------------------------------- cover letter ----

def letter_parts(master: dict, body: str, company: str, role: str) -> dict:
    from datetime import date

    return {
        "name": master["name"],
        "contact": master["contact"],
        "date": date.today().strftime("%B %d, %Y"),
        "re_line": f"Re: {role} — {company}",
        "greeting": "Dear Hiring Manager,",
        "body": " ".join(body.split()),  # enforce single paragraph
        "closing": "Sincerely,",
        "signature": master["name"].title(),
    }


def build_letter_docx(parts: dict, path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(72)
        section.left_margin = section.right_margin = Pt(72)

    style = doc.styles["Normal"]
    style.font.name = "Cambria"
    style.font.size = Pt(11)

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(parts["name"])
    name_run.bold = True
    name_run.font.size = Pt(16)
    doc.add_paragraph(parts["contact"])
    doc.add_paragraph(parts["date"])
    doc.add_paragraph().add_run(parts["re_line"]).bold = True
    doc.add_paragraph(parts["greeting"])
    body_p = doc.add_paragraph(parts["body"])
    body_p.paragraph_format.space_before = Pt(6)
    body_p.paragraph_format.space_after = Pt(12)
    doc.add_paragraph(parts["closing"])
    doc.add_paragraph(parts["signature"])
    doc.save(path)


def build_letter_pdf(parts: dict, path: Path) -> None:
    body_style = ParagraphStyle(
        "letterBody", fontName="Times-Roman", fontSize=11, leading=15
    )
    bold = ParagraphStyle("letterBold", parent=body_style, fontName="Times-Bold")
    name_style = ParagraphStyle(
        "letterName", parent=body_style, fontName="Times-Bold", fontSize=16, leading=20
    )

    def esc(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [
        Paragraph(esc(parts["name"]), name_style),
        Paragraph(esc(parts["contact"]), body_style),
        Spacer(1, 14),
        Paragraph(esc(parts["date"]), body_style),
        Spacer(1, 14),
        Paragraph(esc(parts["re_line"]), bold),
        Spacer(1, 14),
        Paragraph(esc(parts["greeting"]), body_style),
        Spacer(1, 8),
        Paragraph(esc(parts["body"]), body_style),
        Spacer(1, 16),
        Paragraph(esc(parts["closing"]), body_style),
        Spacer(1, 4),
        Paragraph(esc(parts["signature"]), body_style),
    ]
    SimpleDocTemplate(
        str(path), pagesize=letter,
        leftMargin=1 * inch, rightMargin=1 * inch,
        topMargin=1 * inch, bottomMargin=1 * inch,
    ).build(story)

