"""Turn an uploaded resume file into the master-resume JSON the tailor expects.

Conversion is the one step where a model touches the candidate's history before
any check exists to hold it, so the model is told to copy text verbatim, and
the result is checked against the uploaded file: any word in the JSON that the
file never contained is surfaced for the person to review before tailoring.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

from resume_checks import _collect_text, _content_stems, check_master_schema
from tailormade import parse_model_json, run_model

SUPPORTED_SUFFIXES = (".docx", ".pdf", ".txt", ".md", ".json")

EXAMPLE_MASTER = (Path(__file__).parent / "resume_master.example.json").read_text(
    encoding="utf-8"
)

IMPORT_PROMPT = """\
Convert the resume below into JSON with exactly this schema (example values):

{example}

HARD RULES:
1. Copy text VERBATIM from the resume. Do not reword, summarize, correct,
   embellish, or add anything. Every word in your output must appear in the
   resume text.
2. One entry in "bullets" per bullet point in the resume, in the resume's order.
   Do not merge, split, or drop bullets.
3. "org_line" is the employer, location, and dates as the resume states them,
   formatted "Employer, Location | Start - End".
4. "experience" keeps the resume's job order.
5. A section the resume does not have becomes "" or [] - never invent one.
   "tagline" is the headline under the name if there is one, otherwise "".

Return ONLY {{"resume": {{...}}}} - no markdown fences, no commentary.

RESUME TEXT:
{text}
"""


class ResumeImportError(ValueError):
    """The upload could not be read as a resume."""


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
        return "\n".join(line for line in lines if line.strip())
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ModuleNotFoundError:
            raise ResumeImportError("PDF upload needs pypdf: pip install pypdf")
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    raise ResumeImportError(
        f"Unsupported file type {suffix or '(none)'}; use {', '.join(SUPPORTED_SUFFIXES)}"
    )


def import_resume(
    filename: str, data: bytes, log=print, on_model_call=None
) -> tuple[dict, list[str]]:
    """Return (master resume, warnings for the person to review)."""
    if Path(filename).suffix.lower() == ".json":
        try:
            master = json.loads(data.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ResumeImportError(f"Not valid JSON: {exc}")
        return master, check_master_schema(master)

    text = extract_text(filename, data)
    if len(text.split()) < 30:
        raise ResumeImportError(
            "Could not find enough text in that file. A scanned PDF has no text "
            "layer; upload a DOCX or text version instead."
        )
    log(f"Read {len(text.split())} words from {filename}; converting with DeepSeek...")
    master = parse_model_json(
        run_model(IMPORT_PROMPT.format(example=EXAMPLE_MASTER, text=text), on_model_call)
    )["resume"]

    warnings = check_master_schema(master)
    novel = sorted(_content_stems(_collect_text(master)) - _content_stems(text))
    if novel:
        warnings.append(
            f"Conversion used wording not found in your file - check before tailoring: {novel}"
        )
    return master, warnings
